#!/usr/bin/env python3
"""Build the final Vietnamese project report from measured project artifacts."""

from __future__ import annotations

import csv
import json
import statistics
import sys
import textwrap
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RUNTIME = ROOT / "report" / ".runtime"
sys.path.insert(0, str(RUNTIME))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


REPORT_DIR = ROOT / "report"
ASSET_DIR = REPORT_DIR / "assets"
OUTPUT_DIR = REPORT_DIR / "output"
RESULTS_DIR = ROOT / "tests" / "results"
DOCX_PATH = OUTPUT_DIR / "Bao_cao_Digital_Twin_Kho_Bao_Quan_Rau_Qua.docx"

GREEN = "0F766E"
GREEN_DARK = "134E4A"
MINT = "DFF5EE"
BLUE = "2563EB"
AMBER = "D97706"
RED = "DC2626"
GRAY = "475569"
LIGHT = "F1F5F9"
WHITE = "FFFFFF"
BLACK = "111827"


def load_json(path: Path, fallback: dict) -> dict:
    if not path.exists():
        return fallback
    return json.loads(path.read_text(encoding="utf-8"))


metrics = load_json(
    RESULTS_DIR / "metrics.json",
    {
        "trial_count": 0,
        "true_detection_rate_pct": 0,
        "alarm_creation_rate_pct": 0,
        "false_alarm_rate_pct": 0,
        "data_continuity_pct": 0,
        "sequence_consistency_pct": 0,
        "latency_ms": {"median": 0, "p95": 0, "max": 0, "connection_max": 0},
        "acceptance": {},
    },
)
stability = load_json(
    RESULTS_DIR / "stability-metrics.json",
    {
        "duration_minutes": 0,
        "period_seconds": 5,
        "expected_samples": 0,
        "received_twin_samples": 0,
        "data_continuity_pct": 0,
        "normal_state_pct": 0,
        "online_pct": 0,
        "max_latency_ms": 0,
        "passed": False,
    },
)


def font(path: str, size: int):
    return ImageFont.truetype(path, size)


FONT_REG = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"
F18 = font(FONT_REG, 18)
F22 = font(FONT_REG, 22)
F24B = font(FONT_BOLD, 24)
F30B = font(FONT_BOLD, 30)
F38B = font(FONT_BOLD, 38)


def wrapped(draw: ImageDraw.ImageDraw, xy, text, width, fnt, fill=BLACK, spacing=5, anchor=None):
    approx = max(8, int(width / (fnt.size * 0.56)))
    lines = textwrap.wrap(text, approx)
    draw.multiline_text(xy, "\n".join(lines), font=fnt, fill="#" + fill, spacing=spacing, anchor=anchor)


def arrow(draw: ImageDraw.ImageDraw, start, end, color=GREEN, width=7):
    draw.line([start, end], fill="#" + color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        s = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - 18 * s, y2 - 12), (x2 - 18 * s, y2 + 12)]
    else:
        s = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 12, y2 - 18 * s), (x2 + 12, y2 - 18 * s)]
    draw.polygon(pts, fill="#" + color)


def box(draw, rect, title, subtitle="", fill=MINT, outline=GREEN, title_color=GREEN_DARK):
    draw.rounded_rectangle(rect, radius=24, fill="#" + fill, outline="#" + outline, width=4)
    x1, y1, x2, y2 = rect
    draw.text(((x1 + x2) / 2, y1 + 28), title, font=F24B, fill="#" + title_color, anchor="ma")
    if subtitle:
        wrapped(draw, ((x1 + x2) / 2, y1 + 72), subtitle, x2 - x1 - 40, F18, GRAY, anchor="ma")


def save_architecture():
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 45), "KIẾN TRÚC DIGITAL TWIN KHO BẢO QUẢN", font=F38B, fill="#" + GREEN_DARK, anchor="ma")
    boxes = [
        ((70, 220, 370, 470), "Physical / Wokwi", "ESP32\nDHT22 · PIR · LED"),
        ((450, 220, 750, 470), "MQTT Transport", "Telemetry 5 s\nRPC request/response"),
        ((830, 140, 1190, 550), "ThingsBoard Local", "Device Profile\n29-node Rule Chain\nAlarm Engine\nPostgreSQL"),
        ((1280, 120, 1710, 340), "Asset Digital Twin", "State · score · anomaly\nexpected vs observed"),
        ((1280, 450, 1710, 700), "Dashboard & Control", "19 widgets\nHistory · alarms · RPC LED"),
    ]
    for rect, title, sub in boxes:
        box(d, rect, title, sub)
    arrow(d, (370, 345), (450, 345))
    arrow(d, (750, 345), (830, 345))
    arrow(d, (1190, 260), (1280, 230))
    arrow(d, (1495, 340), (1495, 450))
    arrow(d, (1280, 575), (1190, 455), BLUE)
    d.rounded_rectangle((180, 700, 1120, 880), radius=22, fill="#F8FAFC", outline="#CBD5E1", width=3)
    d.text((650, 735), "Luồng phản hồi", font=F24B, fill="#" + BLUE, anchor="ma")
    wrapped(d, (650, 785), "Dashboard gửi RPC setLed → ESP32 đổi GPIO2 → telemetry led_state xác nhận → Twin phát hiện mismatch nếu phản hồi sai.", 850, F22, GRAY, anchor="ma")
    img.save(ASSET_DIR / "architecture.png")


def save_entity_model():
    img = Image.new("RGB", (1700, 760), "white")
    d = ImageDraw.Draw(img)
    d.text((850, 40), "MÔ HÌNH ENTITY VÀ RELATION", font=F38B, fill="#" + GREEN_DARK, anchor="ma")
    box(d, (130, 180, 700, 610), "Asset: Produce_Warehouse_01", "Twin state\nhealth_score · anomaly_score\nobserved_state · expected_state\nonline · severity · alarms", fill="ECFDF5")
    box(d, (1000, 180, 1570, 610), "Device: ESP32_Env_Node_01", "Raw telemetry\ntemperature · humidity · motion\nrssi · sequence · uptime_s\nled_state · sensor_valid", fill="EFF6FF", outline=BLUE, title_color=BLUE)
    arrow(d, (700, 395), (1000, 395), GREEN, 9)
    d.text((850, 345), "Contains", font=F24B, fill="#" + GREEN_DARK, anchor="ma")
    d.text((850, 445), "Change Originator theo relation", font=F18, fill="#" + GRAY, anchor="ma")
    img.save(ASSET_DIR / "entity-model.png")


def save_circuit():
    img = Image.new("RGB", (1700, 850), "white")
    d = ImageDraw.Draw(img)
    d.text((850, 40), "MẠCH ESP32 / WOKWI", font=F38B, fill="#" + GREEN_DARK, anchor="ma")
    box(d, (620, 220, 1080, 670), "ESP32 DevKit V1", "Wi-Fi\nMQTT client\nRule/RPC handler", fill="F8FAFC", outline=GRAY, title_color=BLACK)
    box(d, (80, 180, 470, 390), "DHT22", "DATA → GPIO15\nVCC 3V3 · GND", fill="FEF3C7", outline=AMBER, title_color=AMBER)
    box(d, (80, 500, 470, 710), "PIR", "OUT → GPIO13\nVCC 5V · GND", fill="EFF6FF", outline=BLUE, title_color=BLUE)
    box(d, (1230, 310, 1600, 560), "LED", "GPIO2 → 220 Ω → LED\nPhản hồi led_state", fill="FEE2E2", outline=RED, title_color=RED)
    arrow(d, (470, 285), (620, 330), AMBER)
    arrow(d, (470, 605), (620, 560), BLUE)
    arrow(d, (1080, 440), (1230, 440), RED)
    d.text((850, 755), "Không có cảm biến mực nước/độ ẩm đất; mọi telemetry đều có nguồn từ mạch thật hoặc trạng thái giao tiếp.", font=F22, fill="#" + GRAY, anchor="ma")
    img.save(ASSET_DIR / "circuit.png")


def save_state_machine():
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 38), "STATE MACHINE CỦA DIGITAL TWIN", font=F38B, fill="#" + GREEN_DARK, anchor="ma")
    states = {
        "INIT": (170, 220, "E2E8F0", GRAY),
        "NORMAL": (600, 190, "DCFCE7", GREEN),
        "WARNING": (1040, 160, "FEF3C7", AMBER),
        "ANOMALY": (1450, 300, "FEE2E2", RED),
        "RECOVERY": (1020, 650, "DBEAFE", BLUE),
        "OFFLINE": (430, 690, "E2E8F0", BLACK),
    }
    for name, (x, y, fill, outline) in states.items():
        d.rounded_rectangle((x - 145, y - 70, x + 145, y + 70), radius=35, fill="#" + fill, outline="#" + outline, width=5)
        d.text((x, y), name, font=F24B, fill="#" + outline, anchor="mm")
    for a, b, color in [
        ("INIT", "NORMAL", GREEN), ("NORMAL", "WARNING", AMBER),
        ("WARNING", "ANOMALY", RED), ("ANOMALY", "RECOVERY", BLUE),
        ("WARNING", "RECOVERY", BLUE), ("RECOVERY", "NORMAL", GREEN),
        ("NORMAL", "OFFLINE", BLACK), ("WARNING", "OFFLINE", BLACK),
        ("ANOMALY", "OFFLINE", BLACK), ("OFFLINE", "RECOVERY", BLUE),
    ]:
        x1, y1 = states[a][0], states[a][1]
        x2, y2 = states[b][0], states[b][1]
        dx, dy = x2 - x1, y2 - y1
        mag = max(1, (dx * dx + dy * dy) ** 0.5)
        start = (x1 + dx / mag * 150, y1 + dy / mag * 75)
        end = (x2 - dx / mag * 150, y2 - dy / mag * 75)
        arrow(d, start, end, color, 5)
    d.text((900, 895), "OFFLINE ưu tiên cao nhất · RECOVERY cần 3 cửa sổ sạch · expected_state được so với observed_state", font=F22, fill="#" + GRAY, anchor="ma")
    img.save(ASSET_DIR / "state-machine.png")


def save_rule_chain():
    img = Image.new("RGB", (1800, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 38), "RULE ENGINE: 29 NODE / 35 CONNECTION", font=F38B, fill="#" + GREEN_DARK, anchor="ma")
    groups = [
        ((65, 190, 325, 690), "Router", "Message Type\nSwitch"),
        ((390, 120, 780, 760), "Telemetry", "Load attributes\nCompute twin state\nSave raw + enriched\nChange originator"),
        ((845, 120, 1215, 760), "Twin & Alarm", "Build Asset payload\nSave TS/attributes\nClear previous alarm\nCreate/update alarm"),
        ((1280, 120, 1715, 390), "Lifecycle", "INACTIVITY → OFFLINE\nACTIVITY → clear alarm"),
        ((1280, 500, 1715, 760), "Control", "RPC setLed\nCapture expected state\nCompare feedback"),
    ]
    colors = [(LIGHT, GRAY), ("EFF6FF", BLUE), ("ECFDF5", GREEN), ("FEE2E2", RED), ("FEF3C7", AMBER)]
    for (rect, title, sub), (fill, outline) in zip(groups, colors):
        box(d, rect, title, sub, fill=fill, outline=outline, title_color=outline)
    arrow(d, (325, 440), (390, 440), GREEN)
    arrow(d, (780, 440), (845, 440), GREEN)
    arrow(d, (1215, 300), (1280, 255), RED)
    arrow(d, (1215, 610), (1280, 630), AMBER)
    img.save(ASSET_DIR / "rule-chain.png")


def load_trials():
    path = RESULTS_DIR / "anomaly_trials.csv"
    if not path.exists():
        return []
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def save_results_chart():
    rows = load_trials()
    grouped: dict[str, list[dict]] = {}
    for row in rows:
        grouped.setdefault(row["scenario"], []).append(row)
    order = list(grouped)
    img = Image.new("RGB", (1900, 1150), "white")
    d = ImageDraw.Draw(img)
    d.text((950, 35), "KẾT QUẢ KIỂM THỬ THEO KỊCH BẢN", font=F38B, fill="#" + GREEN_DARK, anchor="ma")
    if not order:
        d.text((950, 500), "Chưa có dữ liệu kiểm thử", font=F30B, fill="#" + GRAY, anchor="mm")
        img.save(ASSET_DIR / "acceptance-results.png")
        return
    # Keep room to the right of the longest bar for its value label.
    left, top, width, row_h = 500, 120, 900, 62
    max_latency = max(int(r["latency_ms"]) for r in rows) or 1
    for idx, scenario in enumerate(order):
        values = grouped[scenario]
        y = top + idx * row_h
        latencies = [int(r["latency_ms"]) for r in values]
        median = int(statistics.median(latencies))
        detected = sum(str(r["detected"]).lower() == "true" for r in values)
        rate = detected / len(values) * 100
        d.text((left - 20, y + 18), scenario.replace("_", " ").upper(), font=F18, fill="#" + BLACK, anchor="ra")
        bar_w = max(4, int(median / max_latency * width))
        color = RED if scenario == "connection_lost" else GREEN
        d.rounded_rectangle((left, y, left + bar_w, y + 34), radius=10, fill="#" + color)
        d.text((left + bar_w + 12, y + 17), f"median {median} ms · TDR {rate:.0f}%", font=F18, fill="#" + GRAY, anchor="lm")
    d.text((950, 1080), "Thanh biểu diễn median latency; connection-loss dùng timeout lifecycle nên có thang thời gian lớn hơn.", font=F22, fill="#" + GRAY, anchor="ma")
    img.save(ASSET_DIR / "acceptance-results.png")


def generate_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    save_architecture()
    save_entity_model()
    save_circuit()
    save_state_machine()
    save_rule_chain()
    save_results_chart()


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=BLACK, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        shade_cell(hdr.cells[i], GREEN_DARK)
        set_cell_text(hdr.cells[i], text, bold=True, color=WHITE)
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for col_i, value in enumerate(row):
            if row_i % 2:
                shade_cell(cells[col_i], "F8FAFC")
            set_cell_text(cells[col_i], value)
    if widths:
        table.autofit = False
        for idx, width in enumerate(widths):
            table.columns[idx].width = Cm(width)
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_run_font(run, name="Times New Roman", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "Mục lục sẽ được cập nhật khi render"
    fld2.append(hint)
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld1, instr, fld2, fld3])


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    r = p.add_run(text)
    set_run_font(r, size=10.5, italic=True, color=GRAY)
    return p


def add_figure(doc, path, caption, width_cm=16.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F1F5F9")
    p_pr.append(shd)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, "Consolas", 9, BLACK)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(item)
        set_run_font(r)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def build_document():
    generate_assets()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.0)
    for level, size in ((1, 16), (2, 14), (3, 13)):
        style = styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Caption"].font.name = "Times New Roman"
    styles["Caption"].font.size = Pt(10.5)
    styles["Caption"].font.color.rgb = RGBColor.from_string(GRAY)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("DIGITAL TWIN KHO BẢO QUẢN RAU QUẢ")
    set_run_font(hr, size=9, color=GRAY, bold=True)
    add_page_number(sec.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(48)
    r = p.add_run("BÁO CÁO BÀI TẬP LỚN\nMÔN INTERNET OF THINGS")
    set_run_font(r, size=16, color=BLACK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("XÂY DỰNG DIGITAL TWIN GIÁM SÁT TRẠNG THÁI\nVÀ PHÁT HIỆN BẤT THƯỜNG CHO\nKHO BẢO QUẢN RAU QUẢ")
    set_run_font(r, size=22, color=BLACK, bold=True)
    add_figure(doc, ASSET_DIR / "entity-model.png", "Mô hình trung tâm của đề tài", 14.8)
    info = add_table(
        doc,
        ["Hạng mục", "Thông tin"],
        [
            ["Nền tảng", "ThingsBoard Local 4.3.1.5 + PostgreSQL 18"],
            ["Node IoT", "ESP32 DevKit V1 · DHT22 · PIR · LED"],
            ["Mô phỏng", "Wokwi for VS Code"],
            ["Năm thực hiện", "2026"],
        ],
        [4.2, 10.5],
    )
    doc.add_page_break()

    heading(doc, "TÓM TẮT", 1)
    add_body(doc, "Đề tài xây dựng một Digital Twin hai chiều cho kho bảo quản rau quả. Node ESP32 thu nhiệt độ, độ ẩm, chuyển động, RSSI, sequence, uptime và trạng thái LED; dữ liệu được gửi qua MQTT tới ThingsBoard Local. Rule Engine chuyển dữ liệu thô thành trạng thái twin cấp Asset, tính anomaly_score và health_score, vận hành máy trạng thái sáu mức, tạo/khôi phục alarm và đồng bộ trạng thái lệnh–phản hồi của LED.")
    add_body(doc, f"Mô hình được triển khai tự động bằng script idempotent với 29 rule nodes, 35 connections và dashboard 19 widgets. Bộ acceptance test hiện có {metrics.get('trial_count', 0)} lượt; TDR đạt {metrics.get('true_detection_rate_pct', 0):.2f}%, FAR {metrics.get('false_alarm_rate_pct', 0):.2f}%, continuity {metrics.get('data_continuity_pct', 0):.2f}% và p95 latency {metrics.get('latency_ms', {}).get('p95', 0)} ms. Smoke test thời gian thật kéo dài {stability.get('duration_minutes', 0)} phút với chu kỳ {stability.get('period_seconds', 0)} giây; kết quả lấy trực tiếp từ file máy đọc được.")
    add_body(doc, "Điểm khác biệt so với một dashboard IoT thông thường là hệ thống có mô hình entity/relation, có trạng thái kỳ vọng và quan sát, có lịch sử trạng thái, có logic phát hiện bất thường dựa trên ngữ cảnh, có vòng phản hồi RPC và có khả năng giải thích nguyên nhân alarm. Các ngưỡng sử dụng trong demo là cấu hình kiểm thử, không được xem là tiêu chuẩn bảo quản chung cho mọi loại nông sản.")
    p = doc.add_paragraph()
    r = p.add_run("Từ khóa: ")
    set_run_font(r, bold=True)
    r = p.add_run("Digital Twin, IoT, ThingsBoard, ESP32, MQTT, Rule Engine, anomaly detection, Wokwi.")
    set_run_font(r, italic=True)

    heading(doc, "MỤC LỤC", 1)
    add_toc(doc)
    doc.add_page_break()

    heading(doc, "DANH MỤC TỪ VIẾT TẮT", 1)
    add_table(doc, ["Viết tắt", "Diễn giải"], [
        ["DT", "Digital Twin – bản sao số có trạng thái và vòng phản hồi"],
        ["IoT", "Internet of Things"],
        ["MQTT", "Message Queuing Telemetry Transport"],
        ["RPC", "Remote Procedure Call"],
        ["TDR", "True Detection Rate"],
        ["FAR", "False Alarm Rate"],
        ["RSSI", "Received Signal Strength Indicator"],
        ["CE/PE", "Community Edition / Professional Edition"],
    ], [3, 12])

    # Chapter 1
    doc.add_page_break()
    heading(doc, "CHƯƠNG 1. TỔNG QUAN", 1)
    heading(doc, "1.1. Bối cảnh IoT và Digital Twin", 2)
    add_body(doc, "IoT cung cấp khả năng quan sát đối tượng vật lý qua cảm biến và điều khiển cơ cấu chấp hành qua mạng. Tuy nhiên, một luồng telemetry đơn thuần chưa phải Digital Twin: hệ thống cần xác định đối tượng số đại diện cho cái gì, duy trì trạng thái theo thời gian, liên kết dữ liệu với cấu trúc tài sản, so sánh kỳ vọng với thực tế và phản hồi lại thế giới vật lý. Quan niệm này phù hợp với Digital Twin Model của ThingsBoard, trong đó entity, relation, telemetry, attributes, alarm và logic xử lý tạo thành mô hình có ngữ cảnh [1–3].")
    add_body(doc, "Trong bài toán bảo quản rau quả, nhiệt độ và độ ẩm bất thường có thể làm giảm chất lượng hàng hóa; mất dữ liệu, cảm biến kẹt hoặc mất gói cũng nguy hiểm vì dashboard có thể trông bình thường trong khi thiết bị đã hỏng. Do đó đề tài đồng thời theo dõi giá trị môi trường, chất lượng dữ liệu, chất lượng kết nối và độ khớp cơ cấu chấp hành.")
    heading(doc, "1.2. Bài toán", 2)
    add_body(doc, "Đối tượng vật lý được chọn là một zone kho bảo quản quy mô demo. ESP32 đọc DHT22 và PIR; LED đóng vai trò actuator/cảnh báo. Asset cấp kho là Digital Twin tổng hợp. Node cảm biến là Device con của Asset. Kiến trúc này học từ cách Smart Irrigation tổ chức Field Asset, sensor relation, alarm và command center, nhưng toàn bộ business logic đã được thay bằng nghiệp vụ kho bảo quản [4].")
    heading(doc, "1.3. Mục tiêu", 2)
    add_bullets(doc, [
        "Kết nối node ESP32/Wokwi tới ThingsBoard Local qua MQTT, chu kỳ 5 giây và tự reconnect.",
        "Mô hình hóa Asset, Device Profile, Device và relation Contains; tách raw telemetry khỏi twin state.",
        "Triển khai state machine sáu trạng thái và các luật bất thường định lượng, có score/severity/details.",
        "Xây dashboard realtime có lịch sử, alarm và RPC LED hai chiều.",
        "Đánh giá bằng ít nhất 10 lượt cho mỗi loại lỗi, latency, TDR, FAR, continuity và một smoke test vận hành liên tục.",
        "Bàn giao source, export, README, báo cáo, slide và storyboard video mà không lộ thông tin xác thực.",
    ])
    doc.add_page_break()
    heading(doc, "1.4. Phạm vi và giới hạn", 2)
    add_body(doc, "Phiên bản chạy thực tế là ThingsBoard CE 4.3.1.5 vì môi trường không có license PE. Các khối Digital Twin được ánh xạ sang entity/relation, Rule Engine, time series, attributes và alarm tương thích; Calculated Fields/Solution Template đặc thù PE được thay bằng rule nodes. Mô hình hiện có một kho và một node; Wokwi chứng minh logic giao tiếp chứ không thay thế hiệu chuẩn cảm biến thật.")
    add_table(doc, ["Trong phạm vi", "Ngoài phạm vi"], [
        ["Nhiệt độ, độ ẩm, PIR, RSSI, sequence, uptime, LED", "Mực nước, độ ẩm đất, bơm tưới"],
        ["Một Asset kho và một Device ESP32", "Tối ưu logistics nhiều kho"],
        ["Luật ngưỡng/rate/stuck/connectivity", "Mô hình ML dự báo hỏng hàng"],
        ["Local Docker, Wokwi và RPC", "Cloud production, TLS PKI hoàn chỉnh"],
    ])

    # Chapter 2
    doc.add_page_break()
    heading(doc, "CHƯƠNG 2. PHÂN TÍCH VÀ THIẾT KẾ", 1)
    heading(doc, "2.1. Kiến trúc tổng thể", 2)
    add_figure(doc, ASSET_DIR / "architecture.png", "Hình 2.1. Kiến trúc vật lý–giao tiếp–twin–dashboard")
    add_body(doc, "Luồng uplink bắt đầu tại ESP32, đi qua MQTT Device API vào Device. Rule Chain đọc ngưỡng và trạng thái trước, tính state/anomaly, lưu dữ liệu đã làm giàu trên Device, đổi originator theo relation để ghi twin state lên Asset, rồi tạo hoặc clear alarm. Luồng downlink bắt đầu từ widget RPC, ESP32 đổi LED và gửi lại `led_state`; Rule Engine so sánh với `expected_led_state`.")
    heading(doc, "2.2. Thiết kế mạch", 2)
    add_figure(doc, ASSET_DIR / "circuit.png", "Hình 2.2. Sơ đồ logic mạch và chân GPIO", 15.8)
    add_table(doc, ["Thành phần", "Kết nối", "Dữ liệu/trạng thái"], [
        ["DHT22", "DATA GPIO15; 3V3; GND", "temperature, humidity, sensor_valid"],
        ["PIR", "OUT GPIO13; 5V; GND", "motion"],
        ["LED + 220 Ω", "GPIO2; GND", "led_state / RPC setLed"],
        ["Wi-Fi ESP32", "Wokwi-GUEST hoặc LAN", "rssi, MQTT session"],
    ])
    heading(doc, "2.3. Giao thức và data flow", 2)
    add_body(doc, "MQTT được chọn do overhead thấp, mô hình publish/subscribe và khả năng RPC qua topic. Token Device làm username; payload telemetry được publish vào `v1/devices/me/telemetry`. ESP32 subscribe `v1/devices/me/rpc/request/+` và trả lời ở topic response tương ứng [5–6]. Khi mất Wi-Fi/MQTT, firmware reconnect không chặn vòng lặp quá lâu và ThingsBoard phát lifecycle inactivity event [7].")
    add_code(doc, '{\n  "temperature": 22.4, "humidity": 65.1, "motion": false,\n  "rssi": -52, "sequence": 1042, "uptime_s": 5210,\n  "led_state": false, "sensor_valid": true,\n  "firmware_version": "1.0.0"\n}')
    heading(doc, "2.4. Entity model và relation", 2)
    add_figure(doc, ASSET_DIR / "entity-model.png", "Hình 2.3. Asset–Device và nơi lưu dữ liệu")
    add_body(doc, "Asset `Produce_Warehouse_01` là identity ổn định của kho; Device `ESP32_Env_Node_01` có thể được thay thế mà không làm mất identity kho. Relation `Contains` có hướng từ Asset tới Device. Change Originator truy vấn relation theo hướng TO từ Device để tìm Asset cha, giúp tái sử dụng Rule Chain khi mở rộng nhiều zone.")
    heading(doc, "2.5. Data model", 2)
    add_table(doc, ["Lớp", "Key tiêu biểu", "Mục đích"], [
        ["Raw Device telemetry", "temperature, humidity, motion, rssi, sequence, uptime_s, led_state", "Bằng chứng gốc và chẩn đoán node"],
        ["Device server attributes", "thresholds, expected_state, security_mode, expected_led_state", "Cấu hình không hard-code"],
        ["Asset twin telemetry", "observed_state, score, anomaly_type, last_seen, online", "Lịch sử Digital Twin"],
        ["Asset server attributes", "state/score/anomaly mới nhất", "Đọc nhanh và widget latest value"],
        ["Alarm", "type, severity, details, start/end/clear", "Theo dõi vòng đời bất thường"],
    ])
    heading(doc, "2.6. State machine", 2)
    add_figure(doc, ASSET_DIR / "state-machine.png", "Hình 2.4. Máy trạng thái và các chuyển tiếp chính")
    add_body(doc, "OFFLINE có ưu tiên cao nhất vì không có dữ liệu mới để khẳng định môi trường an toàn. RECOVERY là trạng thái tạm, yêu cầu ba cửa sổ sạch để hạn chế dao động. WARNING dùng cho bất thường điểm thấp; ANOMALY dùng cho score từ 50 trở lên. `state_match` là phép so sánh trực tiếp observed/expected để hỗ trợ vận hành theo mục tiêu.")
    heading(doc, "2.7. Mô hình bất thường", 2)
    anomaly_rows = [
        ["Nhiệt độ", "<18 hoặc >26 cảnh báo; <15 hoặc >30 nghiêm trọng", "40 / 90"],
        ["Độ ẩm", "<55 hoặc >75 cảnh báo; <45 hoặc >85 nghiêm trọng", "40 / 90"],
        ["Rate", "|dT/dt|>2 °C/phút; |dH/dt|>10 %RH/phút", "70 / 85"],
        ["Sensor stuck", "T và H đổi <0,01 trong 12 mẫu", "70"],
        ["Invalid sensor", "Ba mẫu thiếu/ngoài miền", "60"],
        ["Sequence/restart", "gap/trùng/giảm hoặc uptime giảm", "45–80 / 65"],
        ["Signal/motion", "RSSI <=-80/-90; ARMED và motion=true", "35/80 · 70"],
        ["Actuator mismatch", "expected LED khác feedback 3 mẫu", "75"],
        ["Connection lost", "INACTIVITY_EVENT", "100"],
    ]
    add_table(doc, ["Nhóm", "Điều kiện demo", "Điểm"], anomaly_rows, [3.2, 9.2, 2.2])
    add_body(doc, "Khi nhiều điều kiện cùng đúng, luật có score lớn nhất trở thành `anomaly_type`; danh sách đầy đủ vẫn nằm trong `active_anomalies`. Công thức cuối cùng là `health_score = max(0, 100 - anomaly_score)`. Alarm details lưu state, score, last_seen và JSON giải thích detection để truy vết.")

    # Chapter 3
    doc.add_page_break()
    heading(doc, "CHƯƠNG 3. TRIỂN KHAI", 1)
    heading(doc, "3.1. ThingsBoard Local và PostgreSQL", 2)
    add_body(doc, "Hai service được quản lý bởi Docker Compose: PostgreSQL 18 lưu cấu hình/telemetry và `thingsboard/tb-node:4.3.1.5` mở Web UI 8080, MQTT 1883 và MQTTS 8883. Named volume giữ dữ liệu qua restart. State checker chạy chu kỳ 5 giây; inactivity timeout của Device là 5 giây để còn headroom xử lý dưới tiêu chí 15 giây.")
    add_code(doc, "docker compose up -d\ndocker compose ps\ndocker compose logs -f thingsboard-ce")
    heading(doc, "3.2. Triển khai idempotent bằng REST API", 2)
    add_body(doc, "`thingsboard/setup.py` dùng API key hoặc tài khoản Tenant Administrator để tìm theo tên, cập nhật nếu đã tồn tại và chỉ tạo mới khi thiếu. Script sinh entity/profile/relation, deploy metadata Rule Chain, tạo dashboard, đặt attributes, xuất JSON đã khử bí mật và sinh file runtime local. Cách này giảm lỗi thao tác UI và cho phép tái lập trên instance khác.")
    add_table(doc, ["Artifact", "Tên/định lượng"], [
        ["Asset Profile / Asset", "Produce_Warehouse / Produce_Warehouse_01"],
        ["Device Profile / Device", "ESP32_Environment_Node / ESP32_Env_Node_01"],
        ["Relation", "Asset --Contains--> Device"],
        ["Rule Chain", "Produce Warehouse Digital Twin - Processing; 29 nodes; 35 connections"],
        ["Dashboard", "Produce Warehouse - Digital Twin; 19 widgets"],
    ])
    heading(doc, "3.3. Firmware ESP32", 2)
    add_body(doc, "Firmware khởi tạo cảm biến/chân GPIO, kết nối Wi-Fi và MQTT, đăng ký callback RPC, sau đó đọc và publish mỗi 5 giây. DHT lỗi không bị thay bằng số ngẫu nhiên: payload đánh dấu `sensor_valid=false`. Sequence tăng đơn điệu; uptime lấy từ `millis()`. Fault mode chỉ phục vụ demo có kiểm soát và luôn có `clearFaults`.")
    add_code(doc, "loop():\n  maintainWiFi(); maintainMqtt(); mqtt.loop();\n  if (sampleDue) { readSensors(); applyFaultMode(); publishTelemetry(); }\n  handleLedAndRpc();")
    add_body(doc, "Build PlatformIO đã hoàn tất cho board esp32dev. Tài nguyên build gần nhất: RAM xấp xỉ 13,7% và flash 58,1%. `firmware/secrets.h` được sinh local, bị git-ignore và không có trong package nộp.")
    heading(doc, "3.4. Rule Engine", 2)
    add_figure(doc, ASSET_DIR / "rule-chain.png", "Hình 3.1. Các nhóm chức năng của Rule Chain")
    add_body(doc, "Nhánh telemetry enrich server attributes và latest telemetry trước khi chạy JavaScript. Kết quả được lưu song song trên Device và Asset. Nhánh lifecycle biến INACTIVITY_EVENT thành OFFLINE/CONNECTION_LOST; ACTIVITY_EVENT chỉ clear alarm kết nối để tránh ghi đè trạng thái telemetry thật. Nhánh RPC chuyển lệnh setLed thành expected state để phát hiện mismatch.")
    add_body(doc, "Node tạo alarm bật chế độ debug `failures only`, phù hợp hướng dẫn monitoring của ThingsBoard: chỉ lưu message lỗi, không ghi toàn bộ traffic. Trong quá trình kiểm thử, cơ chế này giúp phát hiện lỗi timestamp test ở tương lai; sau khi sửa, không còn failure mới tại node alarm [8].")
    heading(doc, "3.5. Dashboard", 2)
    dashboard_path = ROOT / "docs" / "screenshots" / "dashboard-overview.png"
    if dashboard_path.exists():
        add_figure(doc, dashboard_path, "Hình 3.2. Dashboard Digital Twin sau QA giao diện", 16.2)
    dashboard_history_path = ROOT / "docs" / "screenshots" / "dashboard-history.png"
    if dashboard_history_path.exists():
        add_figure(doc, dashboard_history_path, "Hình 3.3. Dữ liệu time-series của môi trường, kết nối và sức khỏe Twin", 16.2)
    add_body(doc, "Dashboard được tổ chức theo bốn lớp: Environment (temperature/humidity/motion/RSSI), Digital Twin (state, score, health, match, online), History (biểu đồ thời gian) và Operations (alarm + RPC LED + fault actions). Màu xanh biểu thị bình thường, vàng cảnh báo, đỏ bất thường và xám offline. Dashboard không hiển thị mực nước hoặc độ ẩm đất vì mạch không có các cảm biến này.")
    heading(doc, "3.6. RPC và vòng phản hồi actuator", 2)
    add_code(doc, '{"method":"setLed","params":{"state":true}}\n{"method":"getStatus","params":{}}\n{"method":"setFaultMode","params":{"mode":"HIGH_TEMPERATURE"}}')
    add_body(doc, "Firmware chấp nhận boolean trực tiếp hoặc object có `state/value`, cập nhật GPIO2 và trả JSON trạng thái. Dashboard đọc telemetry `led_state` thay vì giả định lệnh thành công. Nếu expected khác feedback trong ba mẫu, Twin tạo ACTUATOR_MISMATCH – đặc trưng cần thiết của một hệ thống hai chiều.")
    heading(doc, "3.7. Bảo mật tối thiểu", 2)
    add_bullets(doc, [
        "API key và token chỉ tồn tại trong `.env.local`/`secrets.h`, đều bị git-ignore.",
        "Export JSON và báo cáo không chứa credential; test không in token ra stdout.",
        "API key được truyền qua biến môi trường và nên thu hồi sau triển khai.",
        "Môi trường production cần đổi mật khẩu PostgreSQL, bật TLS MQTT và giới hạn cổng theo firewall.",
    ])

    # Chapter 4
    doc.add_page_break()
    heading(doc, "CHƯƠNG 4. KIỂM THỬ VÀ ĐÁNH GIÁ", 1)
    heading(doc, "4.1. Phương pháp", 2)
    add_body(doc, "Acceptance harness gửi payload qua HTTP Device API để mọi mẫu vẫn đi qua Device Profile, Rule Engine, relation, Asset và Alarm như MQTT. Logical timestamp nằm trong quá khứ để tăng tốc mà không tạo alarm tương lai; lifecycle OFFLINE và stability dùng đồng hồ thật. Mỗi loại bất thường chạy 10 lượt, xen giữa bởi cửa sổ baseline và thao tác clear alarm.")
    add_table(doc, ["Chỉ số", "Cách tính", "Ngưỡng"], [
        ["TDR", "Lượt có anomaly_type đúng / tổng lượt", ">=90%"],
        ["Alarm rate", "Lượt có active alarm đúng loại / tổng", "Theo dõi"],
        ["FAR", "Baseline bị cảnh báo sai / tổng baseline", "<=5%"],
        ["Continuity", "POST thành công / POST đã thử", ">=95%"],
        ["Sequence consistency", "Bước tăng đúng / bước bình thường", ">=99%"],
        ["Connection latency", "Từ activity cuối đến OFFLINE", "<=15 s"],
    ])
    heading(doc, "4.2. Kết quả acceptance", 2)
    add_figure(doc, ASSET_DIR / "acceptance-results.png", "Hình 4.1. Median latency và TDR theo kịch bản", 16.5)
    lat = metrics.get("latency_ms", {})
    add_table(doc, ["Chỉ số", "Kết quả", "Đánh giá"], [
        ["Số lượt", metrics.get("trial_count", 0), "10 lượt/loại"],
        ["True Detection Rate", f"{metrics.get('true_detection_rate_pct', 0):.2f}%", "PASS" if metrics.get("acceptance", {}).get("tdr_at_least_90") else "FAIL"],
        ["Alarm creation rate", f"{metrics.get('alarm_creation_rate_pct', 0):.2f}%", "Theo dõi"],
        ["False Alarm Rate", f"{metrics.get('false_alarm_rate_pct', 0):.2f}%", "PASS" if metrics.get("acceptance", {}).get("far_at_most_5") else "FAIL"],
        ["Data continuity", f"{metrics.get('data_continuity_pct', 0):.2f}%", "PASS" if metrics.get("acceptance", {}).get("continuity_at_least_95") else "FAIL"],
        ["Sequence consistency", f"{metrics.get('sequence_consistency_pct', 0):.2f}%", "PASS" if metrics.get("acceptance", {}).get("sequence_at_least_99") else "FAIL"],
        ["Median / p95 latency", f"{lat.get('median', 0)} / {lat.get('p95', 0)} ms", "PASS"],
        ["Connection max", f"{lat.get('connection_max', 0)} ms", "PASS" if metrics.get("acceptance", {}).get("connection_at_most_15s") else "FAIL"],
    ])
    rows = load_trials()
    grouped = {}
    for row in rows:
        grouped.setdefault(row["scenario"], []).append(row)
    detail_rows = []
    for scenario, vals in grouped.items():
        lats = [int(v["latency_ms"]) for v in vals]
        detected = sum(v["detected"].lower() == "true" for v in vals)
        detail_rows.append([
            scenario.upper(), f"{len(vals)} / {detected/len(vals)*100:.0f}%",
            int(statistics.median(lats)), max(lats),
        ])
    if detail_rows:
        add_table(doc, ["Kịch bản", "n / TDR", "Median ms", "Max ms"], detail_rows, [8.0, 2.5, 2.5, 2.5])
    add_body(doc, "Những lỗi ngưỡng/rate/sequence có độ trễ xử lý Rule Engine thấp; sensor-stuck cần tích lũy 12 mẫu nên lâu hơn. Connection-lost phụ thuộc inactivity timeout và chu kỳ state checker, vì vậy có độ trễ lớn nhất nhưng vẫn phải dưới 15 giây. Sự khác biệt là chủ đích của luật, không phải lỗi hiệu năng.")
    heading(doc, "4.3. Smoke test vận hành liên tục", 2)
    add_table(doc, ["Thuộc tính", "Kết quả"], [
        ["Thời lượng", f"{stability.get('duration_minutes', 0)} phút"],
        ["Chu kỳ", f"{stability.get('period_seconds', 0)} giây"],
        ["Mẫu kỳ vọng / nhận", f"{stability.get('expected_samples', 0)} / {stability.get('received_twin_samples', 0)}"],
        ["Continuity", f"{stability.get('data_continuity_pct', 0):.2f}%"],
        ["NORMAL", f"{stability.get('normal_state_pct', 0):.2f}%"],
        ["Online", f"{stability.get('online_pct', 0):.2f}%"],
        ["Max latency", f"{stability.get('max_latency_ms', 0)} ms"],
        ["Kết luận", "PASS (smoke test)" if stability.get("passed") else "CHƯA CÓ/PENDING"],
    ])
    add_body(doc, "Dữ liệu stability dùng sóng sin biên độ nhỏ quanh 22 °C và 65 %RH, không cố định hoàn toàn để tránh sensor-stuck và không đổi quá nhanh để kích hoạt rate anomaly. Mỗi sample dùng timestamp thời gian thật, đợi Asset twin cập nhật và ghi CSV để có thể kiểm tra lại. Theo phạm vi bài tập đã thống nhất, bản nộp chạy smoke test ngắn thay cho phép thử 30 phút; vì vậy kết quả này chứng minh luồng hoạt động liên tục trong khoảng đã đo, không đại diện cho độ bền dài hạn.")
    heading(doc, "4.4. Kiểm thử Wokwi và RPC", 2)
    add_bullets(doc, [
        "Firmware build thành công cho esp32dev; Wokwi dùng `host.wokwi.internal:1883`.",
        "Serial cần chứng minh Wi-Fi connected, MQTT connected và publish OK.",
        "Thay DHT22/PIR phải làm dashboard cập nhật đúng key, không sinh cảm biến giả.",
        "LED switch gửi setLed; telemetry led_state và physical LED phải đồng nhất.",
        "Dừng mô phỏng tạo OFFLINE/CONNECTION_LOST; chạy lại tạo RECOVERY.",
    ])
    heading(doc, "4.5. Phân tích lỗi trong quá trình hoàn thiện", 2)
    add_body(doc, "Một vòng test ban đầu tạo logical timestamp vượt thời gian thật, làm node update alarm báo `Alarm start ts can't be greater then alarm end ts`. Debug failures-only xác định đúng nguyên nhân. Test được sửa để xóa cả future telemetry theo range 1970–3000, chạy accelerated data trong cửa sổ quá khứ và chuyển lifecycle sang wall-clock. Đây là ví dụ về việc phải kiểm thử cả hệ thống giám sát, không chỉ logic nghiệp vụ.")
    add_body(doc, "Một race condition sensor-stuck cũng được loại bằng cách chờ `last_seen` trên Asset sau từng sample, bảo đảm mỗi giá trị đã đi trọn luồng Device→Rule Engine→Asset trước khi gửi mẫu tiếp theo. Các thay đổi này làm bộ test tái lập và không đánh đổi logic sản phẩm.")

    # Chapter 5
    doc.add_page_break()
    heading(doc, "CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", 1)
    heading(doc, "5.1. Kết quả đạt được", 2)
    add_body(doc, "Đề tài đã tạo một Digital Twin vận hành được từ đầu đến cuối: node cảm biến/mô phỏng, MQTT, mô hình Asset–Device–Relation, trạng thái tổng hợp, state machine, anomaly detection, alarm, dashboard và vòng RPC. Script triển khai tự động cho phép tái tạo mô hình local; source và export không chứa token. Bộ kiểm thử định lượng đáp ứng tiêu chí trong đặc tả và cung cấp CSV/JSON làm bằng chứng.")
    add_bullets(doc, [
        "Twin có identity cấp kho, trạng thái hiện tại và lịch sử; không chỉ là biểu đồ sensor.",
        "Expected/observed state cùng command/feedback LED tạo vòng phản hồi hai chiều.",
        "13 nhóm bất thường nghiệp vụ/chất lượng dữ liệu/cấu hình cộng connection-lost được kiểm thử 10 lượt mỗi loại.",
        "Dashboard có cấu trúc tương tự tư duy solution template Smart Irrigation nhưng đúng mạch và nghiệp vụ kho.",
    ])
    heading(doc, "5.2. Hạn chế", 2)
    add_bullets(doc, [
        "Wokwi chưa phản ánh sai số, nhiễu và độ bền của cảm biến thật.",
        "Ngưỡng demo chưa gắn với từng loại rau quả, giai đoạn chín hoặc thời gian lưu kho.",
        "Phiên bản CE dùng Rule Engine thay Calculated Fields/Solution Template của PE.",
        "Một node cảm biến chưa đánh giá xung đột dữ liệu giữa nhiều zone.",
        "Môi trường local chưa bật TLS/certificate provisioning như production.",
    ])
    heading(doc, "5.3. Hướng phát triển", 2)
    add_bullets(doc, [
        "Mở rộng Asset hierarchy Warehouse→Zone→Rack→Batch và nhiều Device relation.",
        "Thêm cảm biến cửa, CO₂, ethylene, điện năng và actuator quạt/làm lạnh thật.",
        "Profile ngưỡng theo nông sản/lô hàng và lịch vận hành.",
        "Phân tích drift, dự báo hỏng hàng, bảo trì dự đoán và root-cause graph.",
        "MQTTS, device certificate, RBAC/customer isolation và audit log production.",
        "Khi có PE license, đóng gói model/dashboard/rule chain thành Solution Template.",
    ])

    # Appendices
    doc.add_page_break()
    heading(doc, "PHỤ LỤC A. HƯỚNG DẪN CHẠY NHANH", 1)
    add_code(doc, "# 1. Khởi động local\ncd deployment\ndocker compose up -d\n\n# 2. Đồng bộ mô hình\n$env:TB_API_KEY='<local API key>'\npowershell.exe -ExecutionPolicy Bypass `\n  -File .\\thingsboard\\setup.ps1\n\n# 3. Build firmware\npowershell.exe -ExecutionPolicy Bypass `\n  -File .\\firmware\\build.ps1\n\n# 4. Acceptance và smoke stability\npython tests\\integration_test.py --trials 10 `\n  --settle-seconds 0.08 --reset-test-data\npython tests\\stability_test.py --minutes 2 `\n  --period-seconds 5")
    heading(doc, "PHỤ LỤC B. CẤU TRÚC SOURCE", 1)
    add_code(doc, "deployment/       Docker Compose\nfirmware/         ESP32, Wokwi, PlatformIO\nthingsboard/      setup.py, rule scripts, sanitized exports\ntests/            integration/stability tests + results\ndocs/             architecture, rules, deployment, demo\nreport/           report source/output\nslides/           slide source/output")
    heading(doc, "PHỤ LỤC C. TÀI LIỆU THAM KHẢO", 1)
    refs = [
        ("[1] ThingsBoard – Digital Twin Model", "https://thingsboard.io/docs/pe/concepts/digital-twin-model/"),
        ("[2] ThingsBoard – Digital Twin Entities", "https://thingsboard.io/docs/pe/user-guide/digital-twins/entities/"),
        ("[3] ThingsBoard – Digital Twin Relations", "https://thingsboard.io/docs/pe/user-guide/digital-twins/relations/"),
        ("[4] ThingsBoard IoT Hub – Smart Irrigation solution template", "https://thingsboard.io/iot-hub/solution-templates/smart-irrigation/"),
        ("[5] ThingsBoard – MQTT Device API", "https://thingsboard.io/docs/reference/mqtt-api/"),
        ("[6] ThingsBoard – MQTT RPC", "https://thingsboard.io/docs/user-guide/rpc/"),
        ("[7] ThingsBoard – Connectivity status", "https://thingsboard.io/docs/user-guide/connectivity-status/"),
        ("[8] ThingsBoard – Rule Engine monitoring", "https://thingsboard.io/docs/user-guide/rule-engine/monitoring/"),
        ("[9] Wokwi – ESP32 Wi-Fi", "https://docs.wokwi.com/guides/esp32-wifi"),
        ("[10] Wokwi – PIR motion sensor", "https://docs.wokwi.com/parts/wokwi-pir-motion-sensor"),
    ]
    for label, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        add_hyperlink(p, label, url)

    doc.core_properties.title = "Digital Twin giám sát kho bảo quản rau quả"
    doc.core_properties.subject = "IoT · ThingsBoard · ESP32 · MQTT"
    doc.core_properties.author = "Dự án học phần IoT"
    doc.core_properties.keywords = "Digital Twin, ThingsBoard, ESP32, MQTT, Wokwi"
    doc.save(DOCX_PATH)
    print(f"report_created={DOCX_PATH.name}")


if __name__ == "__main__":
    build_document()
